#!/usr/bin/env python3
"""将已审读的题目表生成本地学习包；不采集、不生成答案、不改变原始表。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import os
import re
import sys
import uuid
from pathlib import Path
from urllib.parse import urlsplit


ALIASES = {
    "id": ("id", "record_id", "题目ID"),
    "question": ("question", "标准问题", "原文问题", "题目"),
    "topic": ("topic", "主题分类", "主题", "知识点"),
    "chapter": ("chapter", "章节"),
    "level": ("level", "难度"),
    "kind": ("kind", "题目类别"),
    "reference": ("reference", "参考答案", "答题要点"),
    "referenceStatus": ("referenceStatus", "答案状态"),
    "companies": ("companies", "来源公司", "公司"),
    "sources": ("sources",),
    "url": ("url", "来源链接"),
    "title": ("title", "来源标题"),
    "location": ("location", "原文位置"),
    "evidence": ("evidence", "原题依据"),
    "original_ids": ("original_ids", "原始题目"),
}


def text(value, field="文本"):
    if value is None:
        return ""
    if not isinstance(value, str):
        raise ValueError(f"{field} 需要文本，不能将对象或列表猜成题干")
    return value


def array(value, field):
    if value is None or value == "":
        return []
    if isinstance(value, str):
        value = json.loads(value)
    if not isinstance(value, list):
        raise ValueError(f"{field} 需要 JSON 数组")
    return value


def pick(row, key, mapping):
    if key in mapping:
        return row.get(mapping[key])
    return next((row[name] for name in ALIASES[key] if name in row), None)


def read_input(path, sheet=None):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
                raise ValueError("CSV 列名缺失或重复，请先明确列映射")
            rows = list(reader)
            if any(None in row or any(v is None for v in row.values()) for row in rows):
                raise ValueError("CSV 行的列数与表头不一致")
        return rows, {}
    if suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError:
            raise ValueError("Excel 输入需要 openpyxl；按 environment.md 安装 excel 模块") from None
        with path.open("rb") as stream:
            book = openpyxl.load_workbook(stream, read_only=True, data_only=False)
            try:
                if sheet is None and len(book.sheetnames) != 1:
                    raise ValueError("Excel 含多个工作表，请用 --sheet 指定")
                table = book[sheet or book.sheetnames[0]]
                iterator = iter(table.iter_rows(values_only=True))
                headers = list(next(iterator))
                if not all(isinstance(x, str) and x for x in headers) or len(set(headers)) != len(headers):
                    raise ValueError("Excel 列名缺失或重复")
                rows = [dict(zip(headers, row)) for row in iterator if any(x is not None for x in row)]
                if any(isinstance(v, str) and v.startswith("=") for row in rows for v in row.values()):
                    raise ValueError("Excel 含公式；请先确认并导出计算后的值，不能把公式当题目")
            finally:
                book.close()
        return rows, {}
    if suffix == ".ndjson":
        return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()], {}
    if suffix == ".json":
        value = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(value, list):
            return value, {}
        if isinstance(value, dict) and isinstance(value.get("questions"), list):
            return value["questions"], value
        raise ValueError("JSON 应为题目数组或包含 questions 的题包；CLI 信封须先导出为 NDJSON")
    raise ValueError("支持 .csv/.xlsx/.ndjson/.json；Markdown/PDF 请由 Agent 审读后整理为题目数组")


def company_list(value):
    if isinstance(value, str) and not value.startswith("["):
        return [x.strip() for x in re.split(r"[;；、]", value) if x.strip()]
    return [text(x, "公司") for x in array(value, "公司")]


def single_text(value, field):
    if isinstance(value, list):
        if len(value) > 1:
            raise ValueError(f"{field} 有多个选项，请先确定本题主归属")
        value = value[0] if value else None
    return text(value, field)


def safe_url(value):
    value = text(value, "来源链接")
    if value:
        parsed = urlsplit(value)
        if parsed.scheme not in ("http", "https") or not parsed.netloc or parsed.username or parsed.password:
            raise ValueError("来源链接须为不含认证信息的 HTTP(S) 地址；未知链接请留空")
        if any(key in parsed.query.lower() for key in ("xsec_token=", "access_token=", "app_secret=")):
            raise ValueError("来源链接含访问令牌，请使用已核实的规范公开链接")
    return value


def source_item(value):
    if not isinstance(value, dict):
        raise ValueError("sources 每项应为来源对象")
    result = {key: text(value.get(key), key) for key in ("title", "file", "location", "excerpt", "status")}
    result["url"] = safe_url(value.get("url"))
    for key in ("record_id", "table_id"):
        if key in value:
            result[key] = text(value[key], key)
    if "post_ids" in value:
        result["post_ids"] = [text(x, "post_ids") for x in array(value["post_ids"], "post_ids")]
    return result


def linked_ids(value):
    return [text(x.get("id") if isinstance(x, dict) else x, "原题 ID") for x in array(value, "原题关联")]


def original_source(row, table_id):
    return source_item({
        "title": row.get("来源标题"), "url": row.get("来源链接"),
        "location": row.get("原文位置"), "excerpt": row.get("原文问题"),
        "status": "依据已存原题记录；本轮未重新访问原帖", "record_id": row.get("record_id"),
        "table_id": table_id, "post_ids": linked_ids(row.get("所属面经")),
    })


def identity_key(question):
    # 仅复用完全相同的旧条目；题义判断与改写 ID 映射由 Agent 完成。
    return json.dumps([question["question"], sorted(question["companies"]),
                       sorted(s["url"] for s in question["sources"] if s["url"])], ensure_ascii=False)


def build_pack(rows, meta, args):
    mapping = json.loads(args.mapping.read_text(encoding="utf-8")) if args.mapping else {}
    if not isinstance(mapping, dict) or any(k not in ALIASES or not isinstance(v, str) for k, v in mapping.items()):
        raise ValueError("列映射应为规范字段名到实际列名的 JSON 对象")
    previous = json.loads(args.previous.read_text(encoding="utf-8")) if args.previous else None
    pack_id = args.pack_id or meta.get("pack_id")
    if not isinstance(pack_id, str) or not pack_id.strip():
        raise ValueError("须指定稳定的 --pack-id；以后更新沿用同一个 ID")
    if previous and previous.get("pack_id") != pack_id:
        raise ValueError("旧题包 pack_id 不同，不能自动迁移身份")
    originals = {}
    if args.source_records:
        original_rows, _ = read_input(args.source_records)
        for row in original_rows:
            rid = text(row.get("record_id"), "关联原题 record_id")
            if not rid or rid in originals:
                raise ValueError("关联原题记录缺少 ID 或 ID 重复")
            originals[rid] = row
    old_by_content = {}
    for q in previous["questions"] if previous else []:
        old_by_content.setdefault(identity_key(q), []).append(q["id"])
    questions, seen = [], set()
    for index, row in enumerate(rows, 1):
        if not isinstance(row, dict):
            raise ValueError(f"第 {index} 条应为题目对象")
        get = lambda key: pick(row, key, mapping)
        question = text(get("question"), "题目")
        if not question.strip():
            raise ValueError(f"第 {index} 条题干为空；请审读并显式排除，不能静默跳过")
        references = linked_ids(get("original_ids"))
        missing = [rid for rid in references if rid not in originals]
        if missing:
            raise ValueError(f"第 {index} 条尚未读取 {len(missing)} 条关联原题，请提供完整 --source-records")
        if references:
            sources = [original_source(originals[rid], args.source_table_id or "") for rid in dict.fromkeys(references)]
        else:
            sources = [source_item(s) for s in array(get("sources"), "sources")]
        if not sources:
            sources = [source_item({"title": get("title"), "url": get("url"),
                "file": str(args.input), "location": get("location") or f"数据条目 {index}",
                "excerpt": question, "status": "导入文件记录；原帖与面经真实性未自动核验"})]
        q = {"question": question, "companies": company_list(get("companies")), "sources": sources}
        rid = text(get("id"), "题目 ID")
        if rid and not rid.strip():
            raise ValueError("题目 ID 不能只有空白，否则练习进度无法恢复")
        if rid and "record_id" in row and not get("id") == row.get("id"):
            if not args.table_id:
                raise ValueError("飞书 record_id 需要 --table-id 作为身份命名空间")
            rid = f"{args.table_id}:{rid}"
        if not rid:
            matches = old_by_content.get(identity_key(q), [])
            if len(matches) > 1:
                raise ValueError(f"第 {index} 条存在多个同文旧题，请在输入中显式保留 id")
            rid = matches[0] if matches else "q_" + uuid.uuid4().hex
        if rid in seen:
            raise ValueError(f"题目 ID 重复：{rid}；不同子问应各有稳定 ID")
        seen.add(rid)
        topic = single_text(get("topic"), "主题")
        # 飞书单选导出可为单元素数组；多个主题需先选择主归属。
        q.update({"id": rid, "topic": topic or "未分类", "chapter": text(get("chapter")) or topic or "未分类",
            "level": single_text(get("level"), "难度") or "待定", "kind": single_text(get("kind"), "类别") or "导入题目（待核验）",
            "reference": text(get("reference")), "referenceStatus": text(get("referenceStatus")),
            "evidence": text(get("evidence"))})
        if not q["referenceStatus"]:
            q["referenceStatus"] = "输入资料提供，未自动技术校验" if q["reference"] else "尚无参考答案"
        q["revision"] = hashlib.sha256(json.dumps([question, q["reference"]], ensure_ascii=False).encode()).hexdigest()[:24]
        questions.append(q)
    if not questions:
        raise ValueError("输入没有题目")
    topic_input = json.loads(args.topics.read_text(encoding="utf-8")) if args.topics else meta.get("topics", [])
    topic_map = {text(t["name"]): t for t in topic_input}
    if len(topic_map) != len(topic_input):
        raise ValueError("专题名称重复")
    topics = []
    for name in dict.fromkeys(q["topic"] for q in questions):
        info = topic_map.get(name, {})
        topics.append({"name": name, **{key: text(info.get(key)) for key in ("goal", "task", "criteria")}})
    return {"version": 1, "pack_id": pack_id, "title": args.title or text(meta.get("title")) or "面经学习包",
            "description": text(meta.get("description")) or "基于已选材料整理，先作答再核对来源与参考内容。",
            "topics": topics, "questions": questions}


def markdown(pack):
    escape = lambda value: html.escape(value, quote=False)
    lines = ["# " + escape(pack["title"]), "", escape(pack["description"]), ""]
    for topic in pack["topics"]:
        lines.extend(["## " + escape(topic["name"]), ""])
        for key, label in (("goal", "学习目标"), ("task", "实践任务"), ("criteria", "自查标准")):
            if topic[key]:
                lines.extend([f"- {label}：{escape(topic[key])}"])
        lines.append("")
        for q in (q for q in pack["questions"] if q["topic"] == topic["name"]):
            lines.extend(["### " + escape(q["question"]), "", f'{escape(q["level"])} · {escape(q["kind"])} · ID: `{q["id"]}`', ""])
            if q["companies"]:
                lines.extend(["来源公司：" + escape("、".join(q["companies"])), ""])
            if q["evidence"]:
                lines.extend(["归纳依据：" + escape(q["evidence"]), ""])
            for s in q["sources"]:
                lines.extend(["来源：" + escape(" · ".join(x for x in (s["title"], s["file"], s["location"], s["url"]) if x)),
                              "", "> " + escape(s["excerpt"]).replace("\n", "\n> "), "", escape(s["status"]), ""])
            lines.extend(["**我的回答：**", "", "<details><summary>参考内容</summary>", "",
                          escape(q["referenceStatus"]), "", escape(q["reference"]), "", "</details>", ""])
    return "\n".join(lines)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True, help="新的输出目录，不覆盖已有题包")
    parser.add_argument("--pack-id", help="学习库稳定 ID；更新时保持不变")
    parser.add_argument("--title")
    parser.add_argument("--mapping", type=Path, help="规范字段到实际列名的映射")
    parser.add_argument("--topics", type=Path, help="经审读的专题目标、任务、自查说明")
    parser.add_argument("--previous", type=Path, help="上次 pack.json；优先保持输入的 id")
    parser.add_argument("--sheet", help="Excel 工作表名")
    parser.add_argument("--table-id", help="飞书题库表 ID，隔离不同表的记录身份")
    parser.add_argument("--source-records", type=Path, help="关联原题的完整 NDJSON")
    parser.add_argument("--source-table-id", help="关联原题表 ID")
    args = parser.parse_args(argv)
    os.umask(0o077)
    try:
        if args.output_dir.exists():
            raise ValueError("输出目录已存在；请生成新版本目录，避免覆盖题包或个人文件")
        rows, meta = read_input(args.input, args.sheet)
        pack = build_pack(rows, meta, args)
        template = (Path(__file__).resolve().parents[1] / "assets" / "study.html").read_text(encoding="utf-8")
        if template.count("__STUDY_PACK_JSON__") != 1:
            raise ValueError("HTML 模板数据占位符不唯一")
        payload = json.dumps(pack, ensure_ascii=False).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
        rendered = template.replace("__STUDY_PACK_JSON__", payload)
        args.output_dir.mkdir(parents=True)
        (args.output_dir / "pack.json").write_text(json.dumps(pack, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        (args.output_dir / "index.html").write_text(rendered, encoding="utf-8")
        (args.output_dir / "学习包.md").write_text(markdown(pack), encoding="utf-8")
        headers = ["id", "question", "topic", "chapter", "level", "kind", "companies", "reference", "referenceStatus", "evidence", "sources"]
        with (args.output_dir / "题目.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=headers)
            writer.writeheader()
            writer.writerows({key: json.dumps(q[key], ensure_ascii=False) if key in ("companies", "sources") else q[key] for key in headers} for q in pack["questions"])
        print(json.dumps({"output": str(args.output_dir), "pack_id": pack["pack_id"], "questions": len(pack["questions"]),
                          "topics": len(pack["topics"]), "with_reference": sum(bool(q["reference"]) for q in pack["questions"])}, ensure_ascii=False))
        return 0
    except (ValueError, OSError, KeyError, StopIteration) as error:
        print(str(error), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
